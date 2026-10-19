"""把一份 DOCX 的结构完整 dump 到一个 UTF-8 文本文件。

用法：
    <python> resume_dump.py <输入.docx> [输出.txt]
"""

import sys
from pathlib import Path

from docx import Document


def fmt_run(r) -> str:
    bits = []
    if r.bold:
        bits.append("B")
    if r.italic:
        bits.append("I")
    if r.underline:
        bits.append("U")
    if r.font.size:
        bits.append(f"{r.font.size.pt:g}pt")
    if r.font.name:
        bits.append(r.font.name)
    try:
        if r.font.color and r.font.color.rgb:
            bits.append(f"#{r.font.color.rgb}")
    except Exception:
        pass
    return f"<{','.join(bits)}>" if bits else ""


def dump(path: str, out_lines: list):
    doc = Document(path)
    w = out_lines.append

    w(f"=== 文件：{path} ===")
    w(f"段落数 {len(doc.paragraphs)} / 表格数 {len(doc.tables)} / 节数 {len(doc.sections)}")
    w("")

    w("=== 节（页面设置）===")
    for i, s in enumerate(doc.sections):
        w(f"[节 {i}] 页面 {s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm  "
          f"边距 上{s.top_margin.cm:.2f} 下{s.bottom_margin.cm:.2f} "
          f"左{s.left_margin.cm:.2f} 右{s.right_margin.cm:.2f}")
    w("")

    w("=== 默认样式 ===")
    try:
        n = doc.styles["Normal"]
        w(f"Normal: 字体={n.font.name} 字号={n.font.size.pt if n.font.size else None}")
    except Exception as e:
        w(f"（读不到：{e}）")
    w("")

    w("=== 页眉 / 页脚 ===")
    for i, s in enumerate(doc.sections):
        for label, part in (("页眉", s.header), ("页脚", s.footer)):
            texts = [p.text for p in part.paragraphs if p.text.strip()]
            if texts:
                w(f"[节 {i}] {label}: {texts}")
    w("")

    w("=== 正文（按文档实际顺序）===")

    from docx.table import Table
    from docx.text.paragraph import Paragraph

    body = doc.element.body
    idx_p = idx_t = 0
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            idx_p += 1
            if not p.text.strip():
                w(f"[P{idx_p:03d}] (空行)  style={p.style.name}")
                continue
            pf = p.paragraph_format
            info = [f"style={p.style.name}"]
            if pf.space_before is not None:
                info.append(f"段前{pf.space_before.pt:g}")
            if pf.space_after is not None:
                info.append(f"段后{pf.space_after.pt:g}")
            if pf.line_spacing:
                info.append(f"行距{pf.line_spacing}")
            if pf.first_line_indent is not None:
                try:
                    info.append(f"首行{pf.first_line_indent.pt:g}")
                except Exception:
                    pass
            w(f"[P{idx_p:03d}] ({', '.join(info)})")
            w(f"        {p.text!r}")
            for j, r in enumerate(p.runs):
                if r.text:
                    w(f"          run{j}: {r.text!r} {fmt_run(r)}")
        elif tag == "tbl":
            t = Table(child, doc)
            idx_t += 1
            w(f"[T{idx_t:02d}] 表格 {len(t.rows)} 行 x {len(t.columns)} 列  "
              f"style={t.style.name if t.style else None}")
            for ri, row in enumerate(t.rows):
                cells = [c.text.replace("\n", " / ") for c in row.cells]
                w(f"        行{ri}: {cells}")


if __name__ == "__main__":
    lines: list[str] = []
    dump(sys.argv[1], lines)
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("resume_dump.txt")
    out.write_text("\n".join(lines), encoding="utf-8")
    print(f"written: {out}  ({len(lines)} lines)")
